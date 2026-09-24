#!/usr/bin/env python3
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "toolkit-tecnico.docx"
OUT.parent.mkdir(exist_ok=True)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(9)
styles["Normal"].paragraph_format.space_after = Pt(3)
for style_name, size in [("Title", 24), ("Heading 1", 16), ("Heading 2", 12)]:
    st = styles[style_name]
    st.font.name = "Arial"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def borders(table):
    tbl_pr = table._tbl.tblPr
    tbl_borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        tbl_borders.append(el)
    tbl_pr.append(tbl_borders)

def add_table(headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = True
    borders(table)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade(cell, "0E1C33")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx % 2:
                shade(cells[i], "F5F7FA")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(7.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table

title = doc.add_paragraph(style="Title")
title.add_run("Toolkit técnico de mantenimiento")
subtitle = doc.add_paragraph()
subtitle.add_run("Guía de implementación para SAP PM MM CO PP y Power BI").bold = True
doc.add_paragraph("Este documento define cómo desplegar y gobernar la suite. La recomendación central es operar una sola cadena de datos, con cada indicador reconciliado contra su transacción SAP y cada iniciativa de mejora cerrada con ahorro validado y un estándar actualizado.")

doc.add_heading("Decisiones de diseño", level=1)
add_table(["Decisión", "Criterio"], [
    ("Modelo estrella", "Dimensiones filtran hechos con relaciones 1 a N y dirección única."),
    ("Calendario común", "DimFecha gobierna las siete tablas de hechos y evita lógicas temporales duplicadas."),
    ("Trazabilidad SAP", "Cada KPI y visual declara la transacción de origen."),
    ("Brechas explícitas", "El sistema etiqueta proxies y no presenta estimaciones como datos auditados."),
    ("Mejora enfocada", "Las pérdidas se convierten en A3 con criterio de salida financiero y técnico."),
])

doc.add_heading("Arquitectura operativa", level=1)
doc.add_paragraph("La solución separa cinco capas: extracción SAP, transformación Power Query, modelo semántico TMDL, medidas DAX y presentación PBIR. El repositorio mantiene todos los artefactos en texto cuando Power BI lo permite, lo que habilita revisión de cambios y CI.")
add_table(["Capa", "Responsabilidad", "Control principal"], [
    ("Origen", "Extractos PM, MM, CO y PP", "Frecuencia y responsable por transacción"),
    ("Transformación", "Tipos, derivadas y banderas", "Contrato de esquema y reglas reproducibles"),
    ("Modelo", "15 tablas y relaciones", "Claves, cardinalidad y dirección de filtro"),
    ("Semántica", "79 medidas DAX", "Definición, unidad y fuente"),
    ("Presentación", "12 páginas y 174 visuales", "Pregunta de negocio y decisión esperada"),
])

doc.add_heading("Contrato de extracción SAP", level=1)
add_table(["Dominio", "Transacciones", "Uso"], [
    ("Ordenes", "IW39 IW38", "Mix, backlog, cierre, HH y costo por orden"),
    ("Avisos", "IW28 IW29 IW59", "Fallas, causa, duración de parada y MTTR"),
    ("Planes", "IP15 IP24 IP30", "Cumplimiento preventivo y vencimientos"),
    ("Costos", "KOB1 KSB1 S_ALR_87013558", "Real, plan, comprometido y clase de costo"),
    ("Materiales", "MB51 MB52 MM03 ME2N", "Consumo, stock, criticidad y lead time"),
    ("Producción", "COOIS MCRE", "Toneladas, horas de operación y disponibilidad"),
])

doc.add_heading("Definiciones de indicadores", level=1)
add_table(["Indicador", "Definición", "Control"], [
    ("MTTR", "Horas de parada de avisos divididas por número de fallas", "No usar HH de orden"),
    ("MTBF", "Horas equipo operativas divididas por averías", "Multiplicar horas de línea por equipos aplicables"),
    ("Backlog HH", "HH plan menos HH real para órdenes abiertas", "Excluir órdenes cerradas"),
    ("Disponibilidad", "Uno menos horas de detención sobre horas de operación", "Fuente PP más avisos PM"),
    ("Costo por tonelada", "Costo real dividido por toneladas producidas", "Conciliar CO con PP"),
    ("Calidad de datos", "Uno menos banderas sobre controles posibles", "No dividir banderas solo por órdenes"),
])

doc.add_heading("Ritmo de gestión", level=1)
add_table(["Frecuencia", "Foro", "Decisiones"], [
    ("Diaria", "Reunión de mantenimiento", "Emergencias, backlog detenido y repuestos críticos"),
    ("Semanal", "Programación", "Carga, capacidad, PM vencido y cierre técnico"),
    ("Mensual", "Revisión de desempeño", "Costo, confiabilidad y cartera de mejora"),
    ("Trimestral", "Comité de activos", "Prioridades de renovación y cierre de brechas"),
])

doc.add_heading("Ciclo de mejora enfocada", level=1)
for heading, text in [
    ("Plan", "Cuantificar la pérdida, fijar meta y validar causa raíz."),
    ("Do", "Ejecutar contramedidas mediante órdenes PM04 vinculadas al A3."),
    ("Check", "Medir el cambio a 30, 60 y 90 días con la misma definición de KPI."),
    ("Act", "Actualizar el plan preventivo y confirmar el ahorro en CO antes de cerrar."),
]:
    p = doc.add_paragraph()
    p.add_run(f"{heading}. ").bold = True
    p.add_run(text)

doc.add_heading("Calidad de datos y brechas", level=1)
add_table(["Brecha", "Impacto", "Tratamiento"], [
    ("Velocidad nominal", "OEE incompleto", "Integrar MES o maestro de velocidad"),
    ("Defectos y reproceso", "Componente de calidad ausente", "Integrar sistema de calidad"),
    ("Capacidad por puesto", "Backlog en semanas usa proxy", "Mantener CR03 por especialidad"),
    ("Causa raíz", "Pareto poco accionable", "Normalizar catálogos de daño y causa"),
])

doc.add_heading("Despliegue y control de cambios", level=1)
doc.add_paragraph("El equipo debe ejecutar python scripts/build.py y python -m unittest discover -s tests -v antes de cada publicación. GitHub Actions repite el build, valida datos y estructura PBIP, y publica el mockup en GitHub Pages desde la rama main.")
add_table(["Puerta", "Criterio de salida"], [
    ("Datos", "PK y FK válidas, tipos correctos y dataset reproducible"),
    ("Modelo", "15 tablas, al menos 40 medidas y relaciones documentadas"),
    ("Experiencia", "12 páginas, chips SAP y brechas visibles"),
    ("Publicación", "CI verde y artefacto Pages desplegado"),
])

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Maintenance Performance Suite  |  Datos sintéticos  |  Uso metodológico")
footer.runs[0].font.name = "Arial"
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = RGBColor(102, 112, 133)

doc.save(OUT)
print(OUT)
